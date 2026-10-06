from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def export_to_excel(
    file_path,
    sheet_name,
    headers,
    rows
):
    file_path = Path(file_path)

    file_path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = sheet_name

    worksheet.append(headers)

    for cell in worksheet[1]:
        cell.font = Font(
            bold=True
        )

    for row in rows:
        worksheet.append(row)

    for column_cells in worksheet.columns:
        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:
            value = cell.value

            if value is not None:
                max_length = max(
                    max_length,
                    len(str(value))
                )

        worksheet.column_dimensions[
            column_letter
        ].width = max_length + 2

    workbook.save(file_path)

    return file_path